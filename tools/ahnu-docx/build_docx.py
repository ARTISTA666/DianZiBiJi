#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""阶段 4：pandoc body.docx → AHNU 物电格式总装。

结构（学校 Word 模板序 + 研究生院 202003 装订规则）：
  封面（无页码）→ 独创性声明/授权书（无页码）→ [分节·前置 罗马页码]
  摘要 → Abstract → 目录（TOC 字段，1-3 级）→ [分节·正文 阿拉伯页码+页眉]
  第一~七章 → 致谢 → 参考文献 → 附录 A-D
  （现役 PDF 的致谢已随 /tmp 失血丢失、main.tex 仅剩空 clearpage；md 内容为准补回，
   顺序按论文线 424783c 采纳的研究生院口径：正文→致谢→参考文献→附录）

关键转换：
- QQTIKZn/QQIMGn 标记段 → 居中插图（tikz 300dpi PNG 自然宽截 15cm；截图按 LaTeX
  宽度分数 ×15.5cm 版心，trim 按 72dpi=1px/bp 用 PIL 预裁）+ 图下题注（五号宋体居中）
- QQEQtag 标记段与其前显示公式段 → 1×2 无框表（公式居中 + （tag）右端）
- 30 张 pandoc 表 → 表上方题注（表 X-Y　名，五号黑体加粗居中）+ 三线表
  （顶/底 1.5 磅、栏目线 0.75 磅、无竖线、表头加粗重复、五号字）
- 参考文献 118 条 → 五号 + 悬挂缩进 420 twips
"""

import argparse
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

REPO = Path(__file__).resolve().parents[2]
BUILD = REPO / "agent-work/docx-build"

DECL_INDEP = ("本人郑重声明：所呈交的学位论文是本人在导师的指导下独立进行研究工作所取得的成果。"
              "尽我所知，除文中已经注明引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写过的作品成果。"
              "对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式标明。本人完全意识到本声明的法律后果由本人承担。")
DECL_LICENSE = ("本学位论文作者完全了解安徽师范大学有关保留、使用学位论文的规定，即：学校有权保留并向国家有关部门或机构"
                "送交论文的复印件和电子版，允许论文被查阅和借阅；学校可以将学位论文的全部或部分内容编入有关数据库进行检索，"
                "可以采用影印、缩印或扫描等复制手段保存、汇编学位论文。"
                "（保密的学位论文在解密后适用本授权书）")

COVER_ROWS = [
    ("学 科、专业：", "major"),
    ("研 究 方 向：", "research_direction"),
    ("作 者 姓 名：", "author"),
    ("导师及职称：", "supervisor"),
    ("论文提交日期：", "submit_date"),
    ("授予学位日期：", "degree_date"),
]


def set_run(run, text=None, east="宋体", west="Times New Roman", size=12, bold=False):
    if text is not None:
        run.text = text
    run.font.name = west
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), west)
    rf.set(qn("w:hAnsi"), west)
    rf.set(qn("w:eastAsia"), east)
    return run


def para(doc, align=None, before=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_page_number_field(p, size=10.5):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def set_pgnum(section, fmt, start=1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


def cell_bottom_border(cell, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:color"), "000000")
    tb.append(b)
    tcPr.append(tb)


def no_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    tb = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        tb.append(e)
    tblPr.append(tb)


def three_line(tbl):
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tb = OxmlElement("w:tblBorders")
    for edge, sz in (("top", "12"), ("bottom", "12"), ("insideH", None), ("insideV", None),
                     ("left", None), ("right", None)):
        e = OxmlElement(f"w:{edge}")
        if sz is None:
            e.set(qn("w:val"), "none")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), sz)
            e.set(qn("w:color"), "000000")
        tb.append(e)
    tblPr.append(tb)


def tight_cell_margins(tbl):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), "30" if edge in ("top", "bottom") else "60")
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def build_cover(doc, inv):
    c = inv["cover"]
    p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    set_run(p.add_run(), "分类号：" + c["class_no"], east="黑体", size=14)

    for _ in range(3):
        para(doc, after=0)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    set_run(p.add_run(), "安徽师范大学", east="黑体", size=26, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    set_run(p.add_run(), "硕 士 学 位 论 文", east="黑体", size=22, bold=True)

    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    set_run(p.add_run(), c["title_zh"], east="黑体", size=18, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
    set_run(p.add_run(), c["title_en"], size=14, bold=True)

    t = doc.add_table(rows=len(COVER_ROWS), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    for i, (label, key) in enumerate(COVER_ROWS):
        lc, vc = t.rows[i].cells
        lc.width, vc.width = Cm(4.6), Cm(7.6)
        lp = lc.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run(lp.add_run(), label, east="宋体", size=14, bold=True)
        vp = vc.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(vp.add_run(), c[key], east="宋体", size=14)
        cell_bottom_border(vc, 8)

    for _ in range(4):
        para(doc, after=0)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    set_run(p.add_run(), "安徽师范大学硕士学位论文", east="宋体", size=14, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    set_run(p.add_run(), "（二〇二六年九月）", east="宋体", size=14)


def build_declarations(doc):
    page_break(doc)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=18)
    set_run(p.add_run(), "安徽师范大学学位论文独创性声明", east="黑体", size=16, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0)
    set_run(p.add_run(), DECL_INDEP, east="宋体", size=12)
    p.paragraph_format.first_line_indent = Cm(0.85)
    for _ in range(3):
        para(doc, after=0)
    para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0).add_run("")
    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    set_run(p.add_run(), "论文作者签名：　　　　　　　　　签字日期：　　　　年　　月　　日", east="宋体", size=12)

    page_break(doc)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=18)
    set_run(p.add_run(), "安徽师范大学学位论文版权使用授权书", east="黑体", size=16, bold=True)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0)
    set_run(p.add_run(), DECL_LICENSE, east="宋体", size=12)
    p.paragraph_format.first_line_indent = Cm(0.85)
    for _ in range(3):
        para(doc, after=0)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    set_run(p.add_run(), "学位论文作者签名：　　　　　　　　导 师 签 名：　　　　　　　　", east="宋体", size=12)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    set_run(p.add_run(), "签字日期：　　　　年　　月　　日　　　签字日期：　　　　年　　月　　日", east="宋体", size=12)


def build_front_matter(doc, inv):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_pgnum(sec, "upperRoman", 1)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number_field(fp, size=10.5)

    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=18, style="Heading 1")
    set_run(p.add_run(), "摘　要", east="黑体", size=16, bold=True)
    for text in inv["abstract_zh"]:
        p = doc.add_paragraph(style="Body Text")
        set_run(p.add_run(), text, east="宋体", size=12)
    p = para(doc, after=0)
    p.paragraph_format.first_line_indent = Cm(0)
    set_run(p.add_run(), "关键词：", east="黑体", size=12, bold=True)
    set_run(p.add_run(), inv["keywords_zh"], east="宋体", size=12)

    page_break(doc)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=18, style="Heading 1")
    set_run(p.add_run(), "Abstract", east="黑体", size=16, bold=True)
    for text in inv["abstract_en"]:
        p = doc.add_paragraph(style="Body Text")
        set_run(p.add_run(), text, size=12)
    p = para(doc, after=0)
    p.paragraph_format.first_line_indent = Cm(0)
    set_run(p.add_run(), "Key words: ", size=12, bold=True)
    set_run(p.add_run(), inv["keywords_en"], size=12)

    page_break(doc)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    set_run(p.add_run(), "目　录", east="黑体", size=16, bold=True)
    p = para(doc)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "（目录域：在 Word/WPS 中打开时自动更新）"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def build_body_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    set_pgnum(sec, "decimal", 1)
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(hp.add_run(), "安徽师范大学硕士学位论文", east="宋体", size=9)
    pPr = hp._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "000000")
    pb.append(b)
    pPr.append(pb)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number_field(fp, size=10.5)
    return sec


def ptext(p_el):
    return "".join(t.text or "" for t in p_el.findall(".//" + qn("w:t")))


def pstyle(p_el):
    ps = p_el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


def marker_of(text):
    m = re.search(r"^(QQTIKZ\d+QQ|QQIMG\d+QQ|QQEQ[\w-]+QQ)$", text.strip())
    return m.group(1) if m else None


def resolve_image(kind, idx, inv):
    """返回 (图片路径, 宽 cm, 题注)。"""
    from PIL import Image
    if kind == "TIKZ":
        path = BUILD / "tikz" / f"fig{idx:02d}.png"
        W, _ = Image.open(path).size
        return path, min(W * 2.54 / 300.0, 15.0), inv["tikz"][idx - 1]["caption"]
    rec = inv["imgs"][idx - 1]
    src = REPO / "docs" / rec["path"].replace("assets/screenshots/", "user-guide-assets/").replace(
        "assets/innovation-screenshots/", "innovation/screenshots/")
    if not src.exists():
        raise SystemExit(f"image missing: {src}")
    if rec.get("trim_bp"):
        im = Image.open(src)
        W, H = im.size
        l, b, r, t = rec["trim_bp"]  # LaTeX trim 顺序 l b r t；截图 72dpi 下 1px=1bp
        out = BUILD / "tmp" / f"trim{idx:02d}.png"
        out.parent.mkdir(exist_ok=True)
        im.crop((l, t, W - r, H - b)).save(out)
        src = out
    return src, min(rec["width_frac"] * 15.5, 15.0), rec["caption"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(REPO / "docs/毕业论文重构稿-创新点增强版-AHNU物电格式-2026-09-17.docx"))
    args = ap.parse_args()
    inv = json.loads((BUILD / "inventory.json").read_text(encoding="utf-8"))

    doc = Document(str(BUILD / "ahnu_ref.docx"))
    src = Document(str(BUILD / "body.docx"))
    # Pandoc 使用 Compact/VerbatimChar 等样式；缺失的定义会使部分 Word
    # 渲染器把单元格段落解析到表格外。随正文一并迁入缺失样式。
    existing = {s.style_id for s in doc.styles}
    for style in src.styles:
        if style.style_id not in existing:
            doc.styles.element.append(deepcopy(style.element))

    build_cover(doc, inv)
    build_declarations(doc)
    build_front_matter(doc, inv)
    build_body_section(doc)

    body_el = doc.element.body
    sectPr = body_el.find(qn("w:sectPr"))

    def body_append(el):
        # 内容必须位于 sectPr 之前（lxml append 会排到 sectPr 后，Word 判为非法序）
        sectPr.addprevious(el)
        return el

    def body_tail():
        return sectPr.getprevious()

    buckets, order, cur = {}, [], None
    cap_iter = iter(inv["table_captions"])
    pre_h1 = []  # pandoc 在首个标题前输出的 bookmarkStart 等锚点元素，归入首章桶

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def require_cur():
        if cur is None:
            raise SystemExit("元素出现在第一个 H1 之前")

    for child in list(src.element.body):
        if child.tag == qn("w:tbl"):
            require_cur()
            cap = next(cap_iter)
            cp = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2)
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.first_line_indent = Cm(0)
            set_run(cp.add_run(), f"表 {cap['num']}　{cap['name']}", east="黑体", size=10.5, bold=True)
            el = body_append(deepcopy(child))
            restyle_table(Table(el, doc), cap["num"])
            buckets[cur] += [cp._p, el]
            continue
        if child.tag != qn("w:p"):
            if cur is None:
                pre_h1.append(child)
            else:
                buckets[cur].append(body_append(deepcopy(child)))
            continue
        text = ptext(child)
        if pstyle(child) == "Heading1":
            cur = text.strip()
            order.append(cur)
            buckets[cur] = [*(body_append(deepcopy(e)) for e in pre_h1),
                            body_append(deepcopy(child))]
            pre_h1 = []
            continue
        mk = marker_of(text)
        if mk:
            require_cur()
            m = re.match(r"QQ(TIKZ|IMG|EQ)([\w-]+)QQ", mk)
            kind, key = m.group(1), m.group(2)
            if kind == "EQ":
                math_p = body_tail()
                if b"oMathPara" not in etree.tostring(math_p):
                    raise SystemExit(f"式 {key} 前一段不是显示公式段")
                tbl_el = equation_table(doc, math_p, key)
                buckets[cur].remove(math_p)  # 数学段已移入公式表，不得再参与重排
                buckets[cur].append(tbl_el)
            else:
                els = insert_figure(doc, kind, int(key), inv)
                buckets[cur].extend(els)
            continue
        require_cur()
        buckets[cur].append(body_append(deepcopy(child)))

    if next(cap_iter, None) is not None:
        raise SystemExit("题注推断数与实际表格数不一致")

    # ---- 装订序重排：正文章节 → 致谢 → 参考文献 → 附录 ----
    chapters = [k for k in order if k.startswith("第")]
    thanks = [k for k in order if k == "致谢"]
    refs = [k for k in order if k == "参考文献"]
    appendices = [k for k in order if k.startswith("附录")]
    missing = [k for k in order if k not in chapters + thanks + refs + appendices]
    if missing:
        raise SystemExit(f"未分类的 H1 桶: {missing}")
    for k in chapters + thanks + refs + appendices:
        for el in buckets[k]:
            if el.getparent() is not body_el:
                continue  # 已被移入表格内部的元素（公式段）不能再搬动
            body_append(el)  # addprevious 移动节点 → 按装订序重排

    # ---- 参考文献条目样式（五号 + 悬挂缩进 420 twips）----
    for el in buckets[refs[0]]:
        if el.tag != qn("w:p"):
            continue
        p = Paragraph(el, doc)
        pPr = el.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            el.insert(0, pPr)
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:left"), "420")
        ind.set(qn("w:hanging"), "420")
        ind.set(qn("w:firstLine"), "0")
        ind.attrib.pop(qn("w:firstLineChars"), None)
        for r in p.runs:
            set_run(r, east="宋体", size=10.5)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2)

    # ---- 收尾 ----
    for t_el in body_el.findall(".//" + qn("w:t")):
        if t_el.text and "\xa0" in t_el.text:
            t_el.text = t_el.text.replace("\xa0", " ")
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "false")
    doc.settings.element.append(uf)
    populate_toc(doc)
    for name in ("Normal", "Body Text", "First Paragraph", "Block Text", "Compact"):
        if name in doc.styles:
            doc.styles[name].paragraph_format.line_spacing = Pt(20)
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.color.rgb = RGBColor(0, 0, 0)

    doc.core_properties.title = inv["cover"]["title_zh"]
    doc.core_properties.author = ""

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(f"saved -> {args.out}")
    print("buckets:", {k: len(v) for k, v in buckets.items()})



def populate_toc(doc):
    """保留可更新TOC域，并写入已核对页码的可点击缓存，打开即可阅读。"""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    cache = BUILD / "toc-pages.json"
    pages = json.loads(cache.read_text()) if cache.exists() else {}
    entries = [p for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2", "Heading 3")]
    holder = next(p for p in doc.paragraphs if p._p.xpath('.//w:fldSimple[contains(@w:instr,"TOC")]'))
    holder.clear()
    def field(p, kind):
        x = OxmlElement("w:fldChar"); x.set(qn("w:fldCharType"), kind)
        p.add_run()._r.append(x)
    field(holder, "begin")
    instr = OxmlElement("w:instrText"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    holder.add_run()._r.append(instr)
    field(holder, "separate")
    for i, h in enumerate(entries):
        bm = f"ThesisHeading{i}"
        start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"),str(20000+i)); start.set(qn("w:name"),bm)
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"),str(20000+i))
        h._p.insert(1,start); h._p.append(end)
        p = holder if i == 0 else doc.add_paragraph()
        if i: previous.addnext(p._p)
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0); pf.left_indent = Cm((int(h.style.name[-1])-1)*0.5)
        pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = Pt(20)
        pf.tab_stops.add_tab_stop(Cm(15.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"),bm)
        run = p.add_run(h.text); set_run(run,size=10.5,bold=h.style.name=="Heading 1")
        link.append(run._r); p._p.append(link)
        set_run(p.add_run("\t"+pages.get(h.text,"1")),size=10.5)
        previous=p._p
    field(p,"end")


def insert_figure(doc, kind, idx, inv):
    path, w_cm, caption = resolve_image(kind, idx, inv)
    ip = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3)
    ip.paragraph_format.keep_with_next = True
    ip.paragraph_format.line_spacing = 1.0
    ip.paragraph_format.first_line_indent = Cm(0)
    ip.add_run().add_picture(str(path), width=Cm(w_cm))
    cp = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    cp.paragraph_format.first_line_indent = Cm(0)
    set_run(cp.add_run(), caption, east="宋体", size=10.5)
    return [ip._p, cp._p]


def equation_table(doc, math_p_el, tag):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    tight_cell_margins(t)
    lc, rc = t.rows[0].cells
    lc.width, rc.width = Cm(13.0), Cm(2.5)
    lc._tc.append(math_p_el)  # 整段（含 oMathPara）移入左格
    first = lc.paragraphs[0]
    if first._p is not math_p_el and ptext(first._p) == "":
        lc._tc.remove(first._p)
    mp = lc.paragraphs[-1]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp.paragraph_format.line_spacing = 1.0
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(13), Cm(2.5)
    mp.paragraph_format.first_line_indent = Cm(0)
    rc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    rp = rc.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(rp.add_run(), f"（{tag}）", east="宋体", size=12)
    return t._tbl


def restyle_table(tbl, number=""):
    three_line(tbl)
    tight_cell_margins(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 显式写入网格和 tcW；仅有 Pandoc tblGrid 在部分渲染器中会退化为零宽列。
    n = len(tbl.columns)
    widths = [15.5 / n] * n
    if number == "B-2":
        widths = [2.7, 0.8, 3.2, 4.3, 4.5]
    for col, width in zip(tbl.columns, widths):
        col.width = Cm(width)
    tbl.autofit = False
    tw = tbl._tbl.tblPr.find(qn("w:tblW"))
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(Cm(15.5).twips)))
    for ri, row in enumerate(tbl.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if ri == 0:
            repeat_header(row)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            for p in cell.paragraphs:
                p.style = "Normal"
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.first_line_indent = Cm(0)
                for r in p.runs:
                    set_run(r, east="宋体", size=10.5, bold=(ri == 0))
            if ri == 0:
                cell_bottom_border(cell, 6)


if __name__ == "__main__":
    main()
